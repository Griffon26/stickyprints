#!/usr/bin/env python
#
# Stickyprints
# Copyright (C) 2017 Maurice van der Pot <griffon26@kfk4ever.com>
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <http://www.gnu.org/licenses/>.
#

import configparser
import copy
import openpyxl
import os
import re
import tempfile
import xml.etree.ElementTree as ET
import zipfile

from tkinter import *
from tkinter.filedialog import askopenfilename, asksaveasfilename
from tkinter.messagebox import showerror

########################################################################
#
#  Sticky processing code
#
########################################################################
namespaces = {
    'm' : "http://schemas.openxmlformats.org/officeDocument/2006/math",
    'mc' : "http://schemas.openxmlformats.org/markup-compatibility/2006",
    'o' : "urn:schemas-microsoft-com:office:office",
    'r' : "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    'v' : "urn:schemas-microsoft-com:vml",
    'w' : "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    'w10' : "urn:schemas-microsoft-com:office:word",
    'w14' : "http://schemas.microsoft.com/office/word/2010/wordml",
    'wne' : "http://schemas.microsoft.com/office/word/2006/wordml",
    'wp' : "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    'wp14' : "http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing",
    'wpc' : "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas",
    'wpg' : "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
    'wpi' : "http://schemas.microsoft.com/office/word/2010/wordprocessingInk",
    'wps' : "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
}

BODY = '{%s}body' % namespaces['w']
TBL = '{%s}tbl' % namespaces['w']
TBLBORDERS = '{%s}tblBorders' % namespaces['w']
#TRBORDERS = '{%s}trBorders' % namespaces['w']
TCBORDERS = '{%s}tcBorders' % namespaces['w']
BOOKMARKSTART = '{%s}bookmarkStart' % namespaces['w']
BOOKMARKEND = '{%s}bookmarkEnd' % namespaces['w']

def read_tasks_from_excel(excel_filename):
    tasks = []
    wb = openpyxl.load_workbook(excel_filename)
    rows = tuple(wb.worksheets[0].rows)
    header = [str(cell.value) for cell in rows[0]]
    for datarow in rows[1:]:
        data = [str(cell.value) for cell in datarow]
        task = dict(zip(header, data))
        tasks.append(task)
    return tasks

def extract_zipfile_to_dir(filename, dir_to_extract_to):
    with zipfile.ZipFile(filename, 'r') as zf:
        zf.extractall(dir_to_extract_to)

def make_zipfile_from_dir(source_dir, output_filename):
    relroot = os.path.abspath(source_dir)

    with zipfile.ZipFile(output_filename, 'w', zipfile.ZIP_DEFLATED) as zf:
        for root, dirs, files in os.walk(source_dir):
            for filename in files:
                filepath = os.path.join(root, filename)
                if os.path.isfile(filepath):
                    arcname = os.path.join(os.path.relpath(root, relroot), filename)
                    zf.write(filepath, arcname)

def replace_placeholders_with_task_data_in_element(element, task):
        stickytext = ET.tostring(element, encoding='unicode')
        for key, val in task.items():
            stickytext = re.sub('&lt;%s&gt;' % key, val, stickytext, flags = re.IGNORECASE)
        return ET.fromstring(stickytext)

def replace_placeholders_with_task_data(elements, task):
    return [replace_placeholders_with_task_data_in_element(el, task) for el in elements]

# Find the last '<' without a '>', but only if there is no pair of '<>' after it
def find_last_partial_placeholder(element):
    prefix_elem = None

    for elem in element.iter():
        if elem.text:
            lessthan = elem.text.rfind('<')
            if lessthan != -1:
                prefix_elem = None
                if elem.text.find('>', lessthan) == -1:
                    prefix_elem = elem

    return prefix_elem

# Find the first '>', but only if it has no matching '<' and if there is no pair of '<>' before it
def find_first_partial_placeholder(element):
    suffix_elem = None

    for elem in element.iter():
        if elem.text:
            greaterthan = elem.text.find('>')
            if greaterthan != -1:
                if elem.text.rfind('<', 0, greaterthan) == -1:
                    suffix_elem = elem
                break

    return suffix_elem

# This function looks for a '<' in one of the children before the element with
# the specified tag and a '>' in one of the children after the element and 
# glues the content 
def glue_together_broken_placeholders_around_element(parent, element):
    prefix_elem = None
    suffix_elem = None
    elems_to_remove = []

    children = list(parent)

    while children and children[0] is not element:
        child = children.pop(0)
        prefix_elem = find_last_partial_placeholder(child)

    # Skip the element itself
    children.pop(0)

    if prefix_elem != None:
        while children and suffix_elem == None:
            child = children.pop(0)
            elems_to_remove.append(child)
            suffix_elem = find_first_partial_placeholder(child)

    if prefix_elem != None and \
       suffix_elem != None and \
       prefix_elem.tag == suffix_elem.tag:

        prefix_elem.text = prefix_elem.text + suffix_elem.text

        for elem in elems_to_remove:
            parent.remove(elem)

    # Always remove the bookmark
    parent.remove(element)

# The _GoBack bookmark is a bookmark that Word inserts to know where the
# cursor/selection was in the last session. It may be in the middle of a
# placeholder, so that's why we need to remove it.
def remove_go_back_bookmark(rootelement):
    for parent_with_bookmark in rootelement.findall('.//*[%s]' % BOOKMARKSTART):
        bookmark_start = parent_with_bookmark.find("./%s[@{%s}name='_GoBack']" % (BOOKMARKSTART, namespaces['w']))
        if bookmark_start != None:
            bookmark_start_id = bookmark_start.attrib['{%s}id' % namespaces['w']]

            glue_together_broken_placeholders_around_element(parent_with_bookmark, bookmark_start)

            bookmark_end = rootelement.find('.//%s[@{%s}id="%s"]' % (BOOKMARKEND, namespaces['w'], bookmark_start_id))
            if bookmark_end != None:
                glue_together_broken_placeholders_around_element(parent_with_bookmark, bookmark_end)

def remove_all_table_borders(element):
    for borders in element.findall('.//%s' % TBLBORDERS):
        borders.clear()
    for borders in element.findall('.//%s' % TCBORDERS):
        borders.clear()

def create_stickies_from_template(template_filename, tasks, stickies_filename):

    with tempfile.TemporaryDirectory(prefix='stickiestmp') as tempdir:

        extract_zipfile_to_dir(template_filename, tempdir)

        docpath = os.path.join(tempdir, 'word', 'document.xml')
        with open(docpath, 'r') as f:
            xml_content = f.read()

        nr_of_pages = int((len(tasks) - 1) / 6) + 1

        for namespace, urn in namespaces.items():
            ET.register_namespace(namespace, urn)
        tree = ET.fromstring(xml_content)
        body = tree.find(BODY)
        remove_go_back_bookmark(body)
        remove_all_table_borders(body)
        page = copy.deepcopy(list(body))

        # Add as many pages with the same content as the template as we need
        for i in range(nr_of_pages - 1):
            body.extend(page)

        # In each table remove all existing cells and add a copy of the first cell
        template_table_content = list(tree.find('.//%s' % TBL))

        tasks.reverse()
        for table in tree.iter(TBL):
            table.clear()

            if tasks:
                task_table_content = replace_placeholders_with_task_data(template_table_content, tasks.pop())
                table.extend(task_table_content)

        ET.ElementTree(tree).write(docpath, encoding='UTF-8', xml_declaration = True)

        make_zipfile_from_dir(tempdir, stickies_filename)

def generate_stickies(template_filename, tasks_filename, stickies_filename):
    tasks = read_tasks_from_excel(tasks_filename)
    if tasks:
        create_stickies_from_template(template_filename, tasks, stickies_filename)
    else:
        print('No tasks, so no stickies')

########################################################################
#
#  User interface code
#
########################################################################

def get_dirname_from_filename(filename):
    return '.' if not filename else os.path.dirname(filename)

def get_settings_file_path():
    return os.path.expanduser(os.path.join('~', '.stickyprints.conf'))

class MyFrame(Frame):
    def __init__(self):
        Frame.__init__(self)
        self.master.resizable(True, False)
        self.master.minsize(800, 0)
        self.master.title("Stickyprints v0.1")
        self.master.columnconfigure(0, weight=1)
        self.columnconfigure(1, weight=1)
        self.columnconfigure(0, pad=10)
        self.columnconfigure(1, pad=10)
        self.columnconfigure(2, pad=10)
        self.rowconfigure(0, pad=5)
        self.rowconfigure(1, pad=5)
        self.rowconfigure(2, pad=5)
        self.rowconfigure(3, pad=5)
        self.grid(sticky='nsew')

        self.template_label = Label(self, text = 'Template:')
        self.template_entry = Entry(self)
        self.template_button = Button(self, text = 'Change...', command = self.on_change_template_button_clicked)
        self.tasklist_label = Label(self, text = 'Task list (xlsx):')
        self.tasklist_entry = Entry(self)
        self.tasklist_button = Button(self, text = 'Change...', command = self.on_change_tasks_button_clicked)
        self.last_used_stickies_filename = None
        self.generate_button = Button(self, text = 'Generate stickies...', command = self.on_generate_stickies_button_clicked)
        self.status_label = Label(self)
        self.quit_button = Button(self, text = 'Quit', command = self.on_quit_button_clicked)

        self.template_label.grid(row = 0, column = 0)
        self.template_entry.grid(row = 0, column = 1, sticky='ew')
        self.template_button.grid(row = 0, column = 2)
        self.tasklist_label.grid(row = 1, column = 0)
        self.tasklist_entry.grid(row = 1, column = 1, sticky='ew')
        self.tasklist_button.grid(row = 1, column = 2)
        self.generate_button.grid(row = 2, column = 1)
        self.status_label.grid(row = 3, column = 1)
        self.quit_button.grid(row = 3, column = 2)

        self.status_label_default_color = self.status_label.cget('background')

        self.config = configparser.ConfigParser()
        self.load_filenames_from_config()

    def load_filenames_from_config(self):
        self.config.read(get_settings_file_path())

        template_filename = self.config.get('Settings', 'template_file', fallback='')
        self.template_entry.delete(0, END)
        self.template_entry.insert(0, template_filename)

        tasklist_filename = self.config.get('Settings', 'tasklist_file', fallback='')
        self.tasklist_entry.delete(0, END)
        self.tasklist_entry.insert(0, tasklist_filename)

    def save_filenames_to_config(self):
        try:
            self.config.add_section('Settings')
        except configparser.DuplicateSectionError:
            pass
        self.config.set('Settings', 'template_file', self.template_entry.get())
        self.config.set('Settings', 'tasklist_file', self.tasklist_entry.get())
        with open(get_settings_file_path(), 'w') as configfile:
            self.config.write(configfile)

    def on_change_template_button_clicked(self):
        filename = askopenfilename(initialdir=get_dirname_from_filename(self.template_entry.get()),
                                   filetypes=[('Template file', '*.docx')])
        if filename:
            self.template_entry.delete(0, END)
            self.template_entry.insert(0, filename)
            self.save_filenames_to_config()

    def on_change_tasks_button_clicked(self):
        filename = askopenfilename(initialdir=get_dirname_from_filename(self.tasklist_entry.get()),
                                   filetypes=[('Tasks file', '*.xlsx')])
        if filename:
            self.tasklist_entry.delete(0, END)
            self.tasklist_entry.insert(0, filename)
            self.save_filenames_to_config()

    def on_generate_stickies_button_clicked(self):
        if not os.path.exists(self.template_entry.get()):
            showerror('Error', 'The specified template file does not exist. Please select a valid template file.')
            return

        if not os.path.exists(self.tasklist_entry.get()):
            showerror('Error', 'The specified tasklist file does not exist. Please select a valid tasklist file.')
            return

        self.status_label.config(text = 'Generating...', background = self.status_label_default_color)

        try:
            self.save_filenames_to_config()

            filename = asksaveasfilename(initialdir=get_dirname_from_filename(self.last_used_stickies_filename),
                                         filetypes=[('Stickies file', '*.docx')])
            if filename:
                self.last_used_stickies_filename = filename
                generate_stickies(self.template_entry.get(),
                                  self.tasklist_entry.get(),
                                  filename)
                self.status_label.config(text = 'The stickies were generated successfully.', background = "#7FFF00")
            else:
                self.status_label.config(text = 'Cancelled', background = self.status_label_default_color)
        except Exception as e:
            self.status_label.config(text = 'An exception occurred during generation: %s' % e, background = "red")
            raise

    def on_quit_button_clicked(self):
        self.master.destroy()

if __name__ == '__main__':
    MyFrame().mainloop()

